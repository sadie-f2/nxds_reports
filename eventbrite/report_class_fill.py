"""
report_class_fill.py — Upcoming class fill rates for promotion planning.

Shows all live classes sorted by date with registration counts, % full,
and spots remaining. Helps identify which classes need promotion.

Usage:
    EVENTBRITE_MOCK=1 python eventbrite/report_class_fill.py
    EVENTBRITE_MOCK=1 python eventbrite/report_class_fill.py --days 30
    EVENTBRITE_MOCK=1 python eventbrite/report_class_fill.py --output excel --file /tmp/fill.xlsx
    python eventbrite/report_class_fill.py                        # live data from .env
"""

import argparse
import csv
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

# Allow running from repo root or from eventbrite/ dir
sys.path.insert(0, str(Path(__file__).resolve().parent))
import eventbrite


def parse_args():
    p = argparse.ArgumentParser(description="Class fill rate report")
    p.add_argument("--days", type=int, default=60,
                   help="Include classes starting within this many days (default: 60)")
    p.add_argument("--output", choices=["csv", "excel"], default="csv")
    p.add_argument("--file", help="Write to file instead of stdout")
    return p.parse_args()


def build_rows(events: list[dict], cutoff: datetime) -> list[dict]:
    rows = []
    for ev in events:
        start_utc_str = ev.get("start", {}).get("utc", "")
        if not start_utc_str:
            continue

        start_utc = datetime.fromisoformat(start_utc_str.replace("Z", "+00:00"))
        if start_utc > cutoff:
            continue

        ta = ev.get("ticket_availability") or {}
        sold  = ta.get("quantity_sold", 0) or 0
        total = ta.get("quantity_total", 0) or 0
        pct   = round(sold / total * 100) if total else 0
        spots_left = max(total - sold, 0)

        rows.append({
            "class":       ev["name"]["text"],
            "date":        ev["start"]["local"][:10],
            "time":        ev["start"]["local"][11:16],
            "capacity":    total,
            "registered":  sold,
            "pct_full":    pct,
            "spots_left":  spots_left,
            "sold_out":    "Yes" if ta.get("is_sold_out") else "",
            "free":        "Yes" if ev.get("is_free") else "",
            "price":       ta.get("minimum_ticket_price", {}).get("display", ""),
            "url":         ev.get("url", ""),
        })

    rows.sort(key=lambda r: (r["date"], r["time"]))
    return rows


COLUMNS = ["class", "date", "time", "capacity", "registered", "pct_full",
           "spots_left", "sold_out", "free", "price", "url"]


def write_csv(rows: list[dict], file=None):
    out = open(file, "w", newline="") if file else sys.stdout
    try:
        w = csv.DictWriter(out, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
    finally:
        if file:
            out.close()


def write_excel(rows: list[dict], file: str):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        sys.exit("openpyxl not installed — run: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Class Fill Rates"

    # Header row
    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", fgColor="1A1A2E")
    for col, name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=name.replace("_", " ").title())
        cell.font  = header_font
        cell.fill  = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows — highlight low-fill classes (<50%) in amber
    amber_fill = PatternFill("solid", fgColor="FFF3CD")
    for row_idx, row in enumerate(rows, 2):
        for col_idx, key in enumerate(COLUMNS, 1):
            ws.cell(row=row_idx, column=col_idx, value=row[key])
        if row["pct_full"] < 50 and not row["sold_out"]:
            for col_idx in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = amber_fill

    # Column widths
    widths = [40, 12, 8, 10, 12, 10, 12, 10, 6, 10, 50]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    wb.save(file)
    print(f"Saved {len(rows)} rows to {file}", file=sys.stderr)


def main():
    args = parse_args()
    cutoff = datetime.now(timezone.utc) + timedelta(days=args.days)

    print(f"Fetching events…", file=sys.stderr)
    events = eventbrite.fetch_events(status="live")
    print(f"  {len(events)} live events found", file=sys.stderr)

    rows = build_rows(events, cutoff)
    print(f"  {len(rows)} upcoming within {args.days} days", file=sys.stderr)

    if args.output == "excel":
        if not args.file:
            sys.exit("--file required for Excel output")
        write_excel(rows, args.file)
    else:
        write_csv(rows, args.file)
        if args.file:
            print(f"Saved {len(rows)} rows to {args.file}", file=sys.stderr)


if __name__ == "__main__":
    main()
